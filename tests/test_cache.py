"""Comprehensive automated tests for lightweight caching, hit-rate metrics, and immediate invalidation."""

import concurrent.futures
import io
import time
from unittest.mock import patch
import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.api import app
from src.cache import InMemoryCache, cache
from src.database import (
    create_resume,
    delete_resume,
    get_dashboard_summary,
    get_recent_activity,
    get_resume,
    get_resume_model,
    update_resume,
)

client = TestClient(app)


# ---------------------------------------------------------------------------
# 1. InMemoryCache Unit Tests
# ---------------------------------------------------------------------------

def test_cache_set_get_and_miss():
    """Verify basic set, get, and miss behavior in InMemoryCache."""
    test_cache = InMemoryCache(default_ttl=10.0)
    assert test_cache.get("nonexistent") is None
    stats = test_cache.get_stats()
    assert stats["misses"] == 1
    assert stats["hits"] == 0
    assert stats["hit_rate"] == 0.0

    test_cache.set("key1", "value1")
    assert test_cache.get("key1") == "value1"
    stats = test_cache.get_stats()
    assert stats["hits"] == 1
    assert stats["misses"] == 1
    assert stats["hit_rate"] == 0.5


def test_cache_hit_rate_progression():
    """Demonstrate cache hit-rate formula and gradual increase on repeated hits."""
    test_cache = InMemoryCache(default_ttl=10.0)
    test_cache.set("item", 42)

    # 1 miss on another key
    test_cache.get("missing")
    # 4 consecutive hits on item
    for _ in range(4):
        val = test_cache.get("item")
        assert val == 42

    stats = test_cache.get_stats()
    assert stats["hits"] == 4
    assert stats["misses"] == 1
    assert stats["total_requests"] == 5
    assert stats["hit_rate"] == 0.8  # 4 / 5


def test_cache_expiration_ttl():
    """Verify entries expire and are purged after their TTL passes."""
    test_cache = InMemoryCache(default_ttl=0.1)  # 100ms TTL
    test_cache.set("short_lived", "data", ttl=0.05)

    # Immediate get hits
    assert test_cache.get("short_lived") == "data"

    # Sleep past TTL
    time.sleep(0.06)

    # Now expired -> should be a miss
    assert test_cache.get("short_lived") is None
    assert test_cache.get_stats()["size"] == 0


def test_cache_tag_invalidation():
    """Verify invalidation by tag evicts only tagged entries."""
    test_cache = InMemoryCache(default_ttl=60.0)
    test_cache.set("res:1", "data1", tags=["resume", "batch1"])
    test_cache.set("res:2", "data2", tags=["resume", "batch2"])
    test_cache.set("dash:summary", "metrics", tags=["dashboard"])

    assert test_cache.get("res:1") == "data1"
    assert test_cache.get("res:2") == "data2"
    assert test_cache.get("dash:summary") == "metrics"

    # Invalidate by tag 'dashboard'
    count = test_cache.invalidate_by_tag("dashboard")
    assert count == 1
    assert test_cache.get("dash:summary") is None
    assert test_cache.get("res:1") == "data1"
    assert test_cache.get("res:2") == "data2"

    # Invalidate by tag 'batch1'
    count = test_cache.invalidate_by_tag("batch1")
    assert count == 1
    assert test_cache.get("res:1") is None
    assert test_cache.get("res:2") == "data2"


def test_cache_prefix_invalidation():
    """Verify invalidating by key prefix removes matching entries."""
    test_cache = InMemoryCache(default_ttl=60.0)
    test_cache.set("user:1", "alice")
    test_cache.set("user:2", "bob")
    test_cache.set("order:1", "order_data")

    count = test_cache.invalidate_by_prefix("user:")
    assert count == 2
    assert test_cache.get("user:1") is None
    assert test_cache.get("user:2") is None
    assert test_cache.get("order:1") == "order_data"


def test_cache_delete_and_clear():
    """Verify single delete and full clear operations."""
    test_cache = InMemoryCache()
    test_cache.set("k1", "v1")
    test_cache.set("k2", "v2")

    assert test_cache.delete("k1") is True
    assert test_cache.delete("k1") is False
    assert test_cache.get("k1") is None

    test_cache.clear()
    assert test_cache.get("k2") is None
    assert test_cache.get_stats()["size"] == 0


def test_cache_max_size_eviction():
    """Verify cache respects max_size bounds and evicts oldest entries."""
    test_cache = InMemoryCache(default_ttl=60.0, max_size=3)
    test_cache.set("a", 1)
    time.sleep(0.01)
    test_cache.set("b", 2)
    time.sleep(0.01)
    test_cache.set("c", 3)
    time.sleep(0.01)

    # Adding a 4th element should evict the oldest ('a')
    test_cache.set("d", 4)
    assert test_cache.get_stats()["size"] <= 3
    assert test_cache.get("a") is None
    assert test_cache.get("d") == 4


def test_cache_thread_safety():
    """Verify thread-safe concurrent reads and writes without corruption."""
    test_cache = InMemoryCache(default_ttl=10.0, max_size=100)

    def worker(worker_id: int):
        for i in range(50):
            test_cache.set(f"key_{worker_id}_{i}", i, tags=[f"w_{worker_id}"])
            val = test_cache.get(f"key_{worker_id}_{i}")
            assert val == i
            if i % 10 == 0:
                test_cache.invalidate_by_tag(f"w_{worker_id}")

    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as executor:
        futures = [executor.submit(worker, wid) for wid in range(8)]
        for f in concurrent.futures.as_completed(futures):
            f.result()

    stats = test_cache.get_stats()
    assert stats["hits"] > 0
    assert stats["total_requests"] > 0


def test_cache_memoize_decorator():
    """Verify @cache.memoize decorator caches function results."""
    test_cache = InMemoryCache(default_ttl=10.0)
    call_count = 0

    @test_cache.memoize(ttl=5.0, prefix="calc")
    def expensive_calculation(x: int, y: int) -> int:
        nonlocal call_count
        call_count += 1
        return x + y

    # First invocation executes underlying function
    res1 = expensive_calculation(2, 3)
    assert res1 == 5
    assert call_count == 1

    # Second invocation serves from cache
    res2 = expensive_calculation(2, 3)
    assert res2 == 5
    assert call_count == 1

    # Different arguments invoke function
    res3 = expensive_calculation(4, 5)
    assert res3 == 9
    assert call_count == 2


# ---------------------------------------------------------------------------
# 2. Database Integration & Invalidation Tests
# ---------------------------------------------------------------------------

def test_database_get_resume_caching_and_hit_rate():
    """Demonstrate database lookup caching and hit rate increase."""
    cache.clear()
    cache.reset_stats()

    resume_id = create_resume("cached_db_test.pdf", b"test content")

    # Initial get: should populate cache
    r1 = get_resume(resume_id)
    assert r1 is not None
    assert r1["id"] == resume_id

    initial_hits = cache.get_stats()["hits"]

    # Subsequent gets: served from cache
    r2 = get_resume(resume_id)
    assert r2 is not None
    r3 = get_resume_model(resume_id)
    assert r3 is not None

    stats = cache.get_stats()
    assert stats["hits"] > initial_hits
    assert stats["hit_rate"] > 0.0


def test_database_update_resume_immediate_invalidation():
    """Verify update_resume immediately invalidates cached resume and summary."""
    cache.clear()
    resume_id = create_resume("invalidation_test.docx", b"initial")

    # Read into cache
    r_initial = get_resume(resume_id)
    assert r_initial["status"] == "uploaded"
    assert r_initial["user_name"] is None

    # Update resume status and candidate name
    update_resume(resume_id, status="completed", user_name="Alice Candidate")

    # Query again - should immediately reflect updated attributes
    r_updated = get_resume(resume_id)
    assert r_updated["status"] == "completed"
    assert r_updated["user_name"] == "Alice Candidate"


def test_database_delete_resume_immediate_invalidation():
    """Verify delete_resume immediately purges cached record."""
    cache.clear()
    resume_id = create_resume("to_delete.pdf", b"data")
    assert get_resume(resume_id) is not None

    deleted = delete_resume(resume_id)
    assert deleted is True
    assert get_resume(resume_id) is None


def test_database_dashboard_summary_caching_and_invalidation():
    """Verify dashboard summary metrics caching and immediate invalidation on new records."""
    cache.clear()
    cache.reset_stats()

    # Initial summary
    summary_1 = get_dashboard_summary()
    initial_count = summary_1["total_resumes"]

    # Second call hits cache
    summary_2 = get_dashboard_summary()
    assert summary_2["total_resumes"] == initial_count
    assert cache.get_stats()["hits"] >= 1

    # Create new resume - triggers cache invalidation of dashboard summary
    new_id = create_resume("new_candidate.docx", b"resume content")

    # Next call must recompute and reflect new count immediately
    summary_3 = get_dashboard_summary()
    assert summary_3["total_resumes"] == initial_count + 1

    # Cleanup
    delete_resume(new_id)


def test_database_recent_activity_caching():
    """Verify recent activity feed is cached and invalidated on resume creation."""
    cache.clear()
    activity_1 = get_recent_activity(limit=5)
    assert isinstance(activity_1, list)

    new_id = create_resume("recent_activity_check.pdf", b"bytes")
    activity_2 = get_recent_activity(limit=5)
    assert any(item["id"] == new_id for item in activity_2)

    delete_resume(new_id)


# ---------------------------------------------------------------------------
# 3. API Endpoint Caching & Invalidation Tests
# ---------------------------------------------------------------------------

def test_api_dashboard_endpoint_caching():
    """Verify GET /dashboard and GET /dashboard/summary return X-Cache HIT on subsequent calls."""
    cache.clear()

    # First request: Cache MISS
    res1 = client.get("/dashboard/summary")
    assert res1.status_code == 200
    assert res1.headers.get("X-Cache") == "MISS"
    data1 = res1.json()
    assert "total_resumes" in data1
    assert "status_counts" in data1
    assert "recent_activity" in data1

    # Second request: Cache HIT
    res2 = client.get("/dashboard/summary")
    assert res2.status_code == 200
    assert res2.headers.get("X-Cache") == "HIT"
    assert res2.json()["cached"] is True

    # /dashboard route also works identically
    res3 = client.get("/dashboard")
    assert res3.status_code == 200
    assert res3.headers.get("X-Cache") == "HIT"


def test_api_dashboard_immediate_invalidation_on_upload():
    """Verify uploading a new resume immediately invalidates the dashboard cache."""
    cache.clear()

    # Prime dashboard cache
    res1 = client.get("/dashboard/summary")
    assert res1.status_code == 200
    total_before = res1.json()["total_resumes"]

    # Upload new file
    doc = Document()
    doc.add_paragraph("New applicant")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    upload_res = client.post(
        "/upload",
        files={"file": ("applicant.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert upload_res.status_code == 201

    # Next dashboard request should be a MISS with updated count
    res2 = client.get("/dashboard/summary")
    assert res2.status_code == 200
    assert res2.headers.get("X-Cache") == "MISS"
    assert res2.json()["total_resumes"] == total_before + 1

    # Subsequent request is cached HIT
    res3 = client.get("/dashboard/summary")
    assert res3.status_code == 200
    assert res3.headers.get("X-Cache") == "HIT"


def test_api_resume_status_caching_and_invalidation():
    """Verify GET /resumes/{id} demonstrates cache HIT and invalidation on update."""
    cache.clear()
    resume_id = create_resume("status_cache.docx", b"resume data")

    # First request: MISS
    res1 = client.get(f"/resumes/{resume_id}")
    assert res1.status_code == 200
    assert res1.headers.get("X-Cache") == "MISS"
    assert res1.json()["status"] == "uploaded"

    # Second request: HIT
    res2 = client.get(f"/resumes/{resume_id}")
    assert res2.status_code == 200
    assert res2.headers.get("X-Cache") == "HIT"

    # Update the resume
    update_resume(resume_id, status="completed", user_name="Jane Doe")

    # Third request: MISS due to invalidation, but immediately reflects new status
    res3 = client.get(f"/resumes/{resume_id}")
    assert res3.status_code == 200
    assert res3.headers.get("X-Cache") == "MISS"
    assert res3.json()["status"] == "completed"
    assert res3.json()["user_name"] == "Jane Doe"

    # Fourth request: HIT with updated data
    res4 = client.get(f"/resumes/{resume_id}")
    assert res4.status_code == 200
    assert res4.headers.get("X-Cache") == "HIT"
    assert res4.json()["status"] == "completed"


def test_api_cache_stats_endpoint():
    """Verify GET /cache/stats returns operational metrics and hit-rate percentage."""
    cache.clear()
    cache.reset_stats()

    # Generate some cache activity
    res_id = create_resume("metrics_sample.docx", b"content")
    client.get(f"/resumes/{res_id}")  # Miss
    client.get(f"/resumes/{res_id}")  # Hit
    client.get(f"/resumes/{res_id}")  # Hit

    response = client.get("/cache/stats")
    assert response.status_code == 200
    data = response.json()
    assert "hits" in data
    assert "misses" in data
    assert "hit_rate" in data
    assert "total_requests" in data
    assert "size" in data
    assert data["hits"] >= 2
    assert data["hit_rate"] > 0.0


@patch("src.api.call_openai")
def test_api_tailor_flow_invalidates_cache(mock_call_openai):
    """Verify tailoring a resume invalidates its cached status and updates dashboard metrics."""
    cache.clear()
    doc = Document()
    doc.add_paragraph("Mark Spencer")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    resume_id = create_resume("mark.docx", docx_bytes)

    # Prime cache for resume and dashboard
    client.get(f"/resumes/{resume_id}")
    client.get("/dashboard/summary")

    mock_call_openai.return_value = {
        "name": "Mark Spencer",
        "email": "mark@example.com",
        "phone": "555-123-4567",
        "location": "Boston, MA",
        "github": "",
        "linkedin": "",
        "portfolio": "",
        "professional_summary": "Experienced engineer.",
        "work_experience": [],
        "projects": [],
        "skills": ["Python"],
        "soft_skills": [],
        "education": [],
    }

    # Tailor resume
    tailor_res = client.post(
        f"/resumes/{resume_id}/tailor",
        json={"job_description": "Looking for Python specialist with FastAPI experience..."},
    )
    assert tailor_res.status_code == 200

    # Resume status should reflect completed immediately
    status_res = client.get(f"/resumes/{resume_id}")
    assert status_res.status_code == 200
    assert status_res.json()["status"] == "completed"
    assert status_res.json()["has_tailored_resume"] is True

    # Dashboard summary should reflect updated completed count
    dash_res = client.get("/dashboard/summary")
    assert dash_res.status_code == 200
    assert dash_res.json()["status_counts"].get("completed", 0) >= 1
